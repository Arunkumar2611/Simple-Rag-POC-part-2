from __future__ import annotations

import csv
import hashlib
import json
import logging
from pathlib import Path
from typing import Iterator

import pandas as pd
from docx import Document as WordDocument
from langchain_core.documents import Document
from pypdf import PdfReader


logger = logging.getLogger(__name__)
_ocr_engine = None

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".csv", ".xlsx", ".xls", ".docx", ".json", ".md", ".markdown"}


def _ignored_paths(root: Path) -> set[str]:
    ignore_file = root / ".ragignore"
    if not ignore_file.exists():
        return set()
    return {
        line.strip().replace("\\", "/").casefold()
        for line in ignore_file.read_text(encoding="utf-8-sig").splitlines()
        if line.strip() and not line.lstrip().startswith("#")
    }


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()

    with path.open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)

    return digest.hexdigest()


def normalize_text(text: str) -> str:
    text = text.replace("\x00", "")
    return "\n".join(
        line.rstrip() for line in text.splitlines()
    ).strip()


def make_document(text: str,path: Path,root: Path,file_hash: str,content_type: str,**extra_metadata) -> Document | None:
    text = normalize_text(text)
    if not text:
        return None

    metadata = {
        "source": str(path.resolve()),
        "relative_source": str(path.relative_to(root)),
        "file_name": path.name,
        "file_type": path.suffix.lower(),
        "file_hash": file_hash,
        "content_type": content_type,
        **extra_metadata,
    }

    return Document(page_content=text, metadata=metadata)


def load_pdf(path: Path, root: Path, file_hash: str) -> Iterator[Document]:
    reader = PdfReader(str(path))
    extracted = 0
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text(extraction_mode="layout") or page.extract_text() or ""
        extraction_engine = "pypdf"
        if len(text.strip()) < 40:
            try:
                native_text = text.strip()
                ocr_text = _ocr_pdf_page(path, page_number - 1)
                if ocr_text:
                    text = f"{text}\n{ocr_text}".strip()
                    extraction_engine = "pypdf+rapidocr" if native_text else "rapidocr"
            except RuntimeError as exc:
                logger.warning("OCR unavailable for %s page %d: %s", path.name, page_number, exc)
        doc = make_document(
            text=text, path=path, root=root, file_hash=file_hash,
            content_type="pdf_page", page_number=page_number,
            page_count=len(reader.pages), extraction_engine=extraction_engine,
        )
        if doc:
            extracted += 1
            yield doc
    if not extracted:
        raise ValueError("PDF contains no extractable text, including OCR results")


def _ocr_pdf_page(path: Path, page_index: int) -> str:
    """OCR a scanned page using local RapidOCR; no external Tesseract install needed."""
    global _ocr_engine
    try:
        import pypdfium2 as pdfium
        from rapidocr import RapidOCR
    except ImportError as exc:
        raise RuntimeError(
            "Scanned PDF requires OCR dependencies. Run: "
            "python -m pip install pypdfium2 rapidocr onnxruntime"
        ) from exc

    if _ocr_engine is None:
        _ocr_engine = RapidOCR()
    pdf = pdfium.PdfDocument(str(path))
    page = pdf[page_index]
    bitmap = page.render(scale=2.5)
    import numpy as np
    image = np.asarray(bitmap.to_pil().convert("RGB"))
    result = _ocr_engine(image)
    if not result:
        return ""
    return "\n".join(
        text.strip() for text, score in zip(result.txts, result.scores)
        if text and score >= 0.35
    )


def load_text(path: Path, root: Path, file_hash: str) -> Iterator[Document]:
    text = path.read_text(encoding="utf-8-sig")
    doc = make_document(text, path, root, file_hash, "markdown" if path.suffix.lower() in {".md", ".markdown"} else "text")
    if doc:
        yield doc


def load_csv(path: Path, root: Path, file_hash: str) -> Iterator[Document]:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)
        if not reader.fieldnames:
            raise ValueError("CSV has no header row")

        for row_number, row in enumerate(reader, start=2):
            fields = [
                f"{key}: {value}"
                for key, value in row.items()
                if key and value is not None and str(value).strip()
            ]

            if not fields:
                continue

            doc = make_document("\n".join(fields), path, root, file_hash, "csv_row", row_number=row_number)
            if doc:
                yield doc


def load_excel(path: Path, root: Path, file_hash: str,) -> Iterator[Document]:
    sheets = pd.read_excel(path, sheet_name=None, dtype=object)

    for sheet_name, frame in sheets.items():
        frame = frame.where(pd.notna(frame), None)
        columns = [str(column) for column in frame.columns]
        for row_number, values in enumerate(frame.itertuples(index=False, name=None), start=2):
            fields = []
            for column, value in zip(columns, values):
                if value is None:
                    continue

                value = str(value).strip()

                if value:
                    fields.append(f"{column}: {value}")

            if not fields:
                continue

            doc = make_document("\n".join(fields), path, root, file_hash, "excel_row", sheet_name=str(sheet_name), row_number=row_number)
            if doc:
                yield doc


def load_word(path: Path,root: Path,file_hash: str) -> Iterator[Document]:
    word = WordDocument(str(path))
    blocks = []

    for paragraph in word.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        style = paragraph.style.name if paragraph.style else ""
        if style.lower().startswith("heading"):
            blocks.append(f"\n## {text}\n")
        else:
            blocks.append(text)

    for table_index, table in enumerate(word.tables, start=1):
        rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]

        if not rows:
            continue

        width = max(len(row) for row in rows)
        rows = [row + [""] * (width - len(row)) for row in rows]

        header = rows[0]
        markdown = [
            "| " + " | ".join(header) + " |",
            "| " + " | ".join(["---"] * width) + " |",
        ]

        markdown.extend(
            "| " + " | ".join(row) + " |"
            for row in rows[1:]
        )

        blocks.append(
            f"\n### Table {table_index}\n"
            + "\n".join(markdown)
        )

    doc = make_document("\n\n".join(blocks),path,root,file_hash,"word_document")
    if doc:
        yield doc


def load_json(path: Path,root: Path,file_hash: str) -> Iterator[Document]:
    with path.open("r", encoding="utf-8-sig") as file:
        data = json.load(file)

    text = json.dumps(data, ensure_ascii=False, indent=2, default=str)
    doc = make_document(text,path,root,file_hash,"json")
    if doc:
        yield doc


def load_all_documents(data_dir: str,) -> list[Document]:
    root = Path(data_dir).resolve()

    if not root.exists() or not root.is_dir():
        raise NotADirectoryError(
            f"Data directory does not exist: {root}"
        )

    documents: list[Document] = []
    failures = []
    ignored_paths = _ignored_paths(root)

    for path in sorted(root.rglob("*")):
        if not path.is_file():
            continue

        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue

        relative_path = path.relative_to(root).as_posix().casefold()
        if relative_path in ignored_paths:
            logger.info("Skipped ignored file: %s", relative_path)
            continue

        try:
            file_hash = sha256_file(path)
            suffix = path.suffix.lower()

            if suffix == ".pdf":
                loaded = load_pdf(
                    path, root, file_hash
                )
            elif suffix in {".txt", ".md", ".markdown"}:
                loaded = load_text(path, root, file_hash)
            elif suffix == ".csv":
                loaded = load_csv(path, root, file_hash)
            elif suffix in {".xlsx", ".xls"}:
                loaded = load_excel(path, root, file_hash)
            elif suffix == ".docx":
                loaded = load_word(path, root, file_hash)
            elif suffix == ".json":
                loaded = load_json(path, root, file_hash)
            else:
                continue

            file_documents = list(loaded)
            if not file_documents:
                logger.warning(
                    "No content extracted from %s", path
                )
                failures.append({
                    "source": str(path),
                    "error": "No content extracted",
                })
                continue

            documents.extend(file_documents)
            logger.info("Loaded %s: %d document(s)",path,len(file_documents))

        except Exception as exc:
            logger.warning("Failed to load %s: %s", path, exc)
            failures.append({"source": str(path),"error": str(exc)})

    logger.info("Ingestion complete: %d documents, %d failures",len(documents),len(failures))

    if failures:
        logger.warning("Skipped %d file(s) with extraction issues", len(failures))

    return documents


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    docs = load_all_documents("data")
    print(f"Loaded {len(docs)} document records")
        
